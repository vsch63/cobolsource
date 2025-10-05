import os
import pandas as pd
from openpyxl import load_workbook

# Benefit keywords to search for
BENEFIT_KEYWORDS = [
    "Group Life Assurance(Approved)",
    "Group Life Assurance(Unapproved)",
    "PHI",
    "PTD",
    "Funeral Benefit",
    "Spouses Cover",
    "Critical Illness Insurance"
]

def extract_freecover_from_file(file_path, file_name):
    results = []
    try:
        wb = load_workbook(filename=file_path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            num_rows = len(rows)

            # Find the 'INSURED BENEFITS' starting index
            insured_index = None
            for i, row in enumerate(rows):
                if row and any(isinstance(cell, str) and "INSURED BENEFITS" in cell.upper() for cell in row):
                    insured_index = i
                    break

            if insured_index is not None:
                for i in range(insured_index + 1, num_rows):
                    row = rows[i]
                    if not row:
                        continue
                    for j, cell in enumerate(row):
                        if isinstance(cell, str) and cell.strip() in BENEFIT_KEYWORDS:
                            benefit_name = cell.strip()
                            free_cover_value = "N/A"

                            # ✅ FIXED: closed the parenthesis here
                            for k in range(i + 1, num_rows):
                                next_row = rows[k]
                                if not next_row:
                                    continue
                                for col_index, col_val in enumerate(next_row):
                                    if isinstance(col_val, str) and "free cover" in col_val.lower():
                                        if col_index + 1 < len(next_row):
                                            value = next_row[col_index + 1]
                                            if value is not None:
                                                free_cover_value = str(value).strip()
                                        break
                                if free_cover_value != "N/A":
                                    break  # Stop after first match

                            results.append({
                                "File Name": file_name,
                                "Benefit Name": benefit_name,
                                "Free Cover Value": free_cover_value
                            })
        return results
    except Exception as e:
        print(f"❌ Error processing {file_name}: {e}")
        return []

def process_folder_for_freecover(folder_path):
    all_results = []
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            if file.startswith("Conversion Scheme analysis") and file.endswith(".xlsx"):
                file_path = os.path.join(root, file)
                results = extract_freecover_from_file(file_path, file)
                all_results.extend(results)

    if all_results:
        df = pd.DataFrame(all_results)
        output_path = os.path.join(folder_path, "Extracted_Freecover_Details.xlsx")
        df.to_excel(output_path, index=False)
        print(f"✅ Free Cover Output saved to: {output_path}")
    else:
        print("⚠️ No Free Cover details found.")

# 🔁 Run the script
process_folder_for_freecover(r"C:\Users\vscha\Downloads\New folder")
