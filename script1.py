import os
import pandas as pd
from openpyxl import load_workbook

# Keywords to look for
BENEFIT_KEYWORDS = [
    "Group Life Assurance(Approved)",
    "Group Life Assurance(Unapproved)",
    "PHI",
    "PTD",
    "Funeral Benefit",
    "Spouses Cover",
    "Critical Illness Insurance"
]

def extract_benefits_from_file(file_path, file_name):
    results = []
    try:
        wb = load_workbook(filename=file_path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            num_rows = len(rows)

            # Find the 'INSURED BENEFITS' starting index
            insured_index = None
            for i, row in enumerate(rows):
                if row and any(isinstance(cell, str) and "INSURED BENEFITS" in cell.upper() for cell in row):
                    insured_index = i
                    break

            if insured_index is not None:
                for i in range(insured_index + 1, num_rows):
                    row = rows[i]
                    if not row:
                        continue
                    for j, cell in enumerate(row):
                        if isinstance(cell, str) and cell.strip() in BENEFIT_KEYWORDS:
                            benefit_name = cell.strip()
                            formula = "N/A"

                            # Search downward from this row for a row that has "Benefit"
                            for k in range(i + 1, num_rows):
                                next_row = rows[k]
                                if not next_row:
                                    continue
                                for col_index, col_val in enumerate(next_row):
                                    if isinstance(col_val, str) and col_val.strip().lower() == "benefit":
                                        # Get the value in the next column if it exists
                                        if col_index + 1 < len(next_row):
                                            value = next_row[col_index + 1]
                                            if value is not None:
                                                formula = str(value).strip()
                                        break
                                if formula != "N/A":
                                    break  # Stop search after first "Benefit" match

                            results.append({
                                "File Name": file_name,
                                "Benefit Name": benefit_name,
                                "Formula": formula
                            })
        return results
    except Exception as e:
        print(f"❌ Error processing {file_name}: {e}")
        return []

def process_folder(folder_path):
    all_results = []
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            if file.startswith("Conversion Scheme analysis") and file.endswith(".xlsx"):
                file_path = os.path.join(root, file)
                results = extract_benefits_from_file(file_path, file)
                all_results.extend(results)

    if all_results:
        df = pd.DataFrame(all_results)
        output_path = os.path.join(folder_path, "Extracted_Benefit_Details.xlsx")
        df.to_excel(output_path, index=False)
        print(f"✅ Output saved to: {output_path}")
    else:
        print("⚠️ No matching benefit details found.")

# 🚀 Run it
process_folder(r"C:\Users\vscha\Downloads\New folder")
